import io
from flask import Blueprint, request, send_file, g
from openpyxl import Workbook, load_workbook
from ..extensions import db
from ..models import TestCase, Tag, CaseVersion
from ..utils import ok, fail, login_required

bp = Blueprint("excel", __name__, url_prefix="/api/projects/<int:pid>/cases")

PRIORITY_MAP = {"P0": 1, "P1": 2, "P2": 3, "P3": 4}


@bp.get("/export")
@login_required
def export_cases(pid):
    cases = TestCase.query.filter_by(project_id=pid).order_by(TestCase.id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "用例"
    ws.append(["用例编号", "标题", "前置条件", "步骤", "预期结果", "优先级", "类型", "标签"])

    for c in cases:
        steps = c.steps or []
        step_text = "\n".join(f"{i+1}. {s.get('step','')}" for i, s in enumerate(steps))
        expect_text = "\n".join(f"{i+1}. {s.get('expect','')}" for i, s in enumerate(steps))
        ws.append([
            c.case_no, c.title, c.precondition or "",
            step_text, expect_text,
            f"P{c.priority - 1}", c.case_type,
            ",".join(t.name for t in c.tags),
        ])

    for col, w in zip("ABCDEFGH", [14, 40, 30, 40, 40, 8, 12, 20]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True,
        download_name=f"cases_project_{pid}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.post("/import")
@login_required
def import_cases(pid):
    f = request.files.get("file")
    if not f:
        return fail("请上传 Excel 文件")
    try:
        wb = load_workbook(f, data_only=True)
    except Exception as e:
        return fail(f"Excel 解析失败: {e}")

    ws = wb.active
    base_count = TestCase.query.filter_by(project_id=pid).count()
    success, errors = 0, []

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        title = row[0]
        precondition = row[1] if len(row) > 1 else ""
        step_text = row[2] if len(row) > 2 else ""
        expect_text = row[3] if len(row) > 3 else ""
        prio_text = (row[4] if len(row) > 4 else "P1") or "P1"
        case_type = (row[5] if len(row) > 5 else "functional") or "functional"
        tag_text = row[6] if len(row) > 6 else ""

        try:
            # 拆多行步骤（形如 "1. xxx\n2. yyy"），也支持单行
            step_lines = [l.strip() for l in str(step_text or "").split("\n") if l.strip()]
            expect_lines = [l.strip() for l in str(expect_text or "").split("\n") if l.strip()]
            steps = []
            for idx in range(max(len(step_lines), len(expect_lines))):
                steps.append({
                    "step": step_lines[idx] if idx < len(step_lines) else "",
                    "expect": expect_lines[idx] if idx < len(expect_lines) else "",
                })

            base_count += 1
            case = TestCase(
                project_id=pid,
                case_no=f"TC-{base_count:05d}",
                title=title,
                precondition=str(precondition or ""),
                steps=steps,
                priority=PRIORITY_MAP.get(str(prio_text).strip().upper(), 2),
                case_type=case_type,
                creator_id=g.user.id,
            )
            for tag_name in [t.strip() for t in str(tag_text or "").split(",") if t.strip()]:
                tag = Tag.query.filter_by(project_id=pid, name=tag_name).first()
                if not tag:
                    tag = Tag(project_id=pid, name=tag_name)
                    db.session.add(tag)
                case.tags.append(tag)
            db.session.add(case)
            db.session.flush()
            db.session.add(CaseVersion(
                case_id=case.id, version=1,
                snapshot={"title": case.title, "steps": case.steps, "priority": case.priority},
                change_log="Excel导入", operator_id=g.user.id))
            success += 1
        except Exception as e:
            errors.append(f"第 {i} 行: {e}")

    db.session.commit()
    return ok({"success": success, "errors": errors})